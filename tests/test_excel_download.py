"""The "download" Excel mode, against real .xlsx files built with openpyxl.

The SharePoint side is the same in-memory fake the other tests use; the workbooks are
real, so what is asserted is what a person would see when they open the file.
"""
import asyncio
import io
import zipfile

import httpx
import openpyxl
import pytest
from fastapi.testclient import TestClient
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

from app import excel_edit
from app.main import app
from app.sharepoint_client import SharePointClient, SharePointConfigError, get_sharepoint_client
from tests.test_sharepoint import FakeGraph, make_settings

BOOK = "Reports/cases.xlsx"


def table_book(*, data_rows=2, blank_last=False, totals=False, stray_below=False, header=("ID", "Name", "Amount")) -> bytes:
    """A workbook whose sheet "Cases" holds a table "UseCases", plus a second sheet "Notes"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws.append(list(header))
    for i in range(data_rows):
        ws.append([f"r{i}", f"n{i}", i * 10])
    if blank_last:
        ws.append([None, None, None])
    if totals:
        ws.append(["Total", None, "=SUM(C2:C3)"])
    last = ws.max_row
    if last > 1:
        ws.cell(row=last if not totals else last - 1, column=3).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 30
    ws.merge_cells("E1:F1")
    if stray_below:
        ws.cell(row=last + 1, column=1, value="stray note")
    ref = f"A1:C{last}"
    tab = Table(displayName="UseCases", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    tab.autoFilter = AutoFilter(ref=ref)
    if totals:
        tab.totalsRowCount = 1
    ws.add_table(tab)
    notes = wb.create_sheet("Notes")
    notes["A1"] = "keep me"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def plain_book(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                if value is not None:
                    ws.cell(row=r, column=c, value=value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def load(graph: FakeGraph, path: str = BOOK) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(graph.data(path)))


@pytest.fixture
def graph() -> FakeGraph:
    return FakeGraph()


@pytest.fixture
def client(graph: FakeGraph) -> SharePointClient:
    return SharePointClient(
        make_settings(sharepoint_excel_mode="download"),
        transport=httpx.MockTransport(graph.graph_handler),
        token_transport=httpx.MockTransport(graph.token_handler),
    )


@pytest.fixture
def api(client: SharePointClient):
    app.dependency_overrides[get_sharepoint_client] = lambda: client
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def post(api, **body):
    body.setdefault("path", "cases.xlsx")
    return api.post("/sharepoint/append_excel_row", json=body)


# ------------------------------------------------------------------ tables


def test_row_is_added_to_the_end_of_the_table(graph, api):
    graph.add_file(BOOK, table_book(data_rows=2))
    resp = post(api, table="UseCases", values=["x", "Ada", 5.5])
    assert resp.status_code == 201
    assert resp.json() == {"path": "cases.xlsx", "mode": "table", "table": "UseCases", "sheet": None, "row_number": 3, "address": None, "notes": []}

    ws = load(graph)["Cases"]
    assert [c.value for c in ws[4]][:3] == ["x", "Ada", 5.5]
    tbl = ws.tables["UseCases"]
    assert tbl.ref == "A1:C4" and tbl.autoFilter.ref == "A1:C4"
    assert ws["C4"].number_format == ws["C3"].number_format == "#,##0.00"  # looks like the rows above it
    assert isinstance(ws["C4"].value, float)


def test_everything_else_in_the_workbook_survives(graph, api):
    graph.add_file(BOOK, table_book(data_rows=2))
    assert post(api, table="UseCases", values=["x", "y", 1]).status_code == 201
    wb = load(graph)
    assert wb["Notes"]["A1"].value == "keep me"
    assert wb["Cases"].column_dimensions["A"].width == 30
    assert "E1:F1" in {str(r) for r in wb["Cases"].merged_cells.ranges}
    assert [c.value for c in wb["Cases"][2]][:3] == ["r0", "n0", 0]  # existing rows untouched


def test_cells_by_column_name_in_any_order(graph, api):
    graph.add_file(BOOK, table_book(data_rows=1))
    assert post(api, table="UseCases", row={"Amount": 7, "Name": "Cy"}).status_code == 201
    cells = [c.value for c in load(graph)["Cases"][3]][:3]
    assert cells[0] in (None, "") and cells[1:] == ["Cy", 7]


def test_table_name_matching_ignores_case(graph, api):
    graph.add_file(BOOK, table_book())
    assert post(api, table="usecases", values=["a", "b", 1]).status_code == 201


def test_wrong_column_names_counts_and_tables_are_rejected_without_touching_the_file(graph, api):
    original = table_book()
    graph.add_file(BOOK, original)
    resp = post(api, table="UseCases", row={"Nme": "typo"})
    assert resp.status_code == 400 and "Name" in resp.json()["detail"]
    resp = post(api, table="UseCases", values=["only one"])
    assert resp.status_code == 400 and "3 columns" in resp.json()["detail"]
    resp = post(api, table="Nope", values=["a", "b", 1])
    assert resp.status_code == 404 and "Nope" in resp.json()["detail"]
    assert graph.data(BOOK) == original


def test_a_blank_row_at_the_end_of_a_table_is_filled_not_skipped(graph, api):
    graph.add_file(BOOK, table_book(data_rows=0, blank_last=True))  # header + one empty row, like a fresh template
    assert post(api, table="UseCases", values=["a", "b", 1]).json()["row_number"] == 1
    ws = load(graph)["Cases"]
    assert ws.tables["UseCases"].ref == "A1:C2" and ws["A2"].value == "a"
    assert post(api, table="UseCases", values=["c", "d", 2]).json()["row_number"] == 2
    ws = load(graph)["Cases"]
    assert ws.tables["UseCases"].ref == "A1:C3" and ws["A3"].value == "c"


def test_data_under_the_table_is_never_overwritten(graph, api):
    original = table_book(data_rows=2, stray_below=True)
    graph.add_file(BOOK, original)
    resp = post(api, table="UseCases", values=["x", "y", 1])
    assert resp.status_code == 409 and "A4" in resp.json()["detail"]
    assert graph.data(BOOK) == original


def test_a_table_with_a_totals_row_is_refused(graph, api):
    original = table_book(data_rows=2, totals=True)
    graph.add_file(BOOK, original)
    resp = post(api, table="UseCases", values=["x", "y", 1])
    assert resp.status_code == 400 and "totals row" in resp.json()["detail"]
    assert graph.data(BOOK) == original


def test_calculated_columns_are_reported_not_silently_left_empty(graph, api):
    wb = openpyxl.load_workbook(io.BytesIO(table_book()))
    wb["Cases"].tables["UseCases"].tableColumns[2].calculatedColumnFormula = TableFormula(attr_text="UseCases[[#This Row],[ID]]")
    out = io.BytesIO()
    wb.save(out)
    graph.add_file(BOOK, out.getvalue())
    notes = post(api, table="UseCases", values=["x", "y", 1]).json()["notes"]
    assert any("Calculated columns" in n and "Amount" in n for n in notes)


def test_text_stays_text_and_formulas_need_permission(graph, api):
    graph.add_file(BOOK, table_book())
    assert post(api, table="UseCases", values=["2026-10-13", "t", True]).status_code == 201
    ws = load(graph)["Cases"]
    assert ws["A4"].value == "2026-10-13" and isinstance(ws["A4"].value, str)
    assert ws["C4"].value is True

    before = graph.data(BOOK)
    assert post(api, table="UseCases", values=["=HYPERLINK(\"http://evil\")", "t", 1]).status_code == 400
    assert graph.data(BOOK) == before
    assert post(api, table="UseCases", values=["=SUM(1,1)", "t", 1], allow_formulas=True).status_code == 201
    assert load(graph)["Cases"]["A5"].value == "=SUM(1,1)"


# ------------------------------------------------------------------ sheets without a table


def test_below_the_data_on_an_empty_sheet_then_the_next_row(graph, api):
    graph.add_file(BOOK, plain_book({"Data": [], "Other": [[None, "h1", "h2"], [None, "a", "b"]]}))
    first = post(api, values=["x", 1, 2]).json()
    assert (first["sheet"], first["address"], first["row_number"]) == ("Data", "A1:C1", 1)
    assert post(api, values=["y", 3, 4]).json()["address"] == "A2:C2"
    ws = load(graph)["Data"]
    assert [c.value for c in ws[2]][:3] == ["y", 3, 4]


def test_follows_the_used_range_of_a_named_sheet(graph, api):
    graph.add_file(BOOK, plain_book({"Data": [], "Other": [[None, "h1", "h2"], [None, "a", "b"], [None, "c", "d"]]}))
    resp = post(api, sheet="other", values=["e", "f"]).json()  # sheet names match regardless of case
    assert (resp["sheet"], resp["address"], resp["row_number"]) == ("Other", "B4:C4", 4)
    assert load(graph)["Data"].max_row == 1


def test_missing_sheet_and_empty_values(graph, api):
    graph.add_file(BOOK, plain_book({"Data": []}))
    resp = post(api, sheet="Third", values=["x"])
    assert resp.status_code == 404 and "Third" in resp.json()["detail"]
    assert post(api, values=[]).status_code == 400


# ------------------------------------------------------------------ not Excel


def test_files_that_are_not_excel_get_415_in_this_mode_too(graph, api):
    graph.add_file("Reports/data.csv", b"a,b\n1,2\n")
    graph.add_file("Reports/renamed.xlsx", b"a,b\n1,2\n")
    graph.add_file("Reports/empty.xlsx", b"")
    zipped = io.BytesIO()
    with zipfile.ZipFile(zipped, "w") as z:  # a real zip, but a Word document, not a workbook
        z.writestr("word/document.xml", "<w:document/>")
    graph.add_file("Reports/word-in-disguise.xlsx", zipped.getvalue())
    for name in ("data.csv", "renamed.xlsx", "empty.xlsx", "word-in-disguise.xlsx"):
        resp = post(api, path=name, table="UseCases", values=["a", "b", 1])
        assert resp.status_code == 415, name
        assert "Excel" in resp.json()["detail"]


# ------------------------------------------------------------------ people editing at the same time


def test_a_change_made_by_someone_else_mid_append_is_kept(graph, api):
    graph.add_file(BOOK, table_book(data_rows=2))
    original_put, calls = graph._put_item, {"n": 0}

    def racing(path, data, conflict, if_match):
        calls["n"] += 1
        if calls["n"] == 1:  # someone saves a new row between our download and our upload
            wb = load(graph)
            ws = wb["Cases"]
            ws.append(["someone-else", "x", 1])
            ws.tables["UseCases"].ref = ws.tables["UseCases"].autoFilter.ref = "A1:C4"
            out = io.BytesIO()
            wb.save(out)
            graph.items[path]["data"], graph.items[path]["etag"] = out.getvalue(), '"changed"'
        return original_put(path, data, conflict, if_match)

    graph._put_item = racing
    resp = post(api, table="UseCases", values=["mine", "y", 2])
    assert resp.status_code == 201 and calls["n"] == 2
    ws = load(graph)["Cases"]
    assert [ws["A4"].value, ws["A5"].value] == ["someone-else", "mine"]  # neither row was lost
    assert resp.json()["row_number"] == 4


def test_it_gives_up_with_409_if_the_file_never_stops_changing(graph, api, monkeypatch):
    async def no_sleep(_):
        return None

    monkeypatch.setattr(asyncio, "sleep", no_sleep)
    graph.add_file(BOOK, table_book())
    original_put = graph._put_item

    def always_racing(path, data, conflict, if_match):
        graph._counter += 1
        graph.items[path]["etag"] = f'"{graph._counter + 1000}"'
        return original_put(path, data, conflict, if_match)

    graph._put_item = always_racing
    resp = post(api, table="UseCases", values=["x", "y", 1])
    assert resp.status_code == 409 and "kept changing" in resp.json()["detail"]


def test_simultaneous_appends_in_one_process_both_land(graph, client):
    graph.add_file(BOOK, table_book(data_rows=2))

    async def go():
        return await asyncio.gather(
            *[
                client.append_excel_row(["cases.xlsx"], values=[f"u{i}", "n", i], row=None, table="UseCases", sheet=None, allow_formulas=False)
                for i in range(3)
            ]
        )

    results = asyncio.run(go())
    assert sorted(r["row_number"] for r in results) == [3, 4, 5]
    ws = load(graph)["Cases"]
    assert {ws.cell(row=r, column=1).value for r in (4, 5, 6)} == {"u0", "u1", "u2"}
    assert ws.tables["UseCases"].ref == "A1:C6"


# ------------------------------------------------------------------ notes and configuration


def test_features_openpyxl_drops_are_reported_in_notes(graph, api, monkeypatch):
    real_load = openpyxl.load_workbook

    def noisy_load(*args, **kwargs):
        import warnings

        warnings.warn("Data Validation extension is not supported and will be removed")
        return real_load(*args, **kwargs)

    monkeypatch.setattr(excel_edit.openpyxl, "load_workbook", noisy_load)
    graph.add_file(BOOK, table_book())
    notes = post(api, table="UseCases", values=["x", "y", 1]).json()["notes"]
    assert notes == ["Data Validation extension is not supported and will be removed"]


def test_an_unknown_mode_is_a_configuration_error():
    with pytest.raises(SharePointConfigError) as exc:
        SharePointClient(make_settings(sharepoint_excel_mode="bogus"))
    assert "download" in str(exc.value) and "graph" in str(exc.value)


def test_download_is_the_default_mode():
    from app.config import Settings

    assert Settings().sharepoint_excel_mode == "download"
